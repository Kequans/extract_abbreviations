#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Kequan
缩略词提取工具
从Word文档中提取格式为"中文（英文全称, 英文缩写）"的缩略词
"""

import re
from docx import Document
import sys


# Normalize typographical variants for matching; retain the author's display spelling.
_DASHES = str.maketrans({c: "-" for c in "‐‑‒–—−"})
_PARENTHESES = re.compile(r"[（(]([^()（）]+)[)）]")
_ABBREVIATION = re.compile(r"[A-Za-z0-9]+(?:[-/+&.][A-Za-z0-9]+)*")
_ENGLISH_NAME = re.compile(r"[A-Za-z0-9][A-Za-z0-9\s\-/&+.'’]*")
# Explicit prose boundaries, rather than deleting individual Chinese characters.
_CONTEXT = re.compile(
    r"以及|也就是|也称为|称为|简称|记为|分别对应|对应|也包括|包括|"
    r"提出了一种|提出一种|提出了|提出|设计了|设计|构建了|构建|"
    r"本文采用|采用|通过控制|通过|利用|引入|借助|基于|结合|扩展|强化|"
    r"封装至|替换为|随着|也称|即为|中如|导致|加大了|由多层"
)


def _clean_space(text):
    return " ".join(text.split())


def _is_abbreviation(text, include_single_letter):
    normalized = text.translate(_DASHES)
    letters = re.sub(r"[^A-Za-z]", "", normalized)
    return bool(
        _ABBREVIATION.fullmatch(normalized)
        and any(c.isupper() for c in letters)
        and (len(letters) >= 2 or (include_single_letter and len(letters) == 1))
    )


def _parse_definition(body, include_single_letter):
    parts = re.split(r"[,，;；]", body)
    if len(parts) != 2:
        return None
    left, right = map(_clean_space, parts)
    for full, abbr in ((left, right), (right, left)):
        normalized = full.translate(_DASHES)
        if not _is_abbreviation(abbr, include_single_letter):
            continue
        if not _ENGLISH_NAME.fullmatch(normalized) or not re.search(r"[A-Za-z]", full):
            continue
        # Reject two acronym-like tokens, numbers and bibliographic citations.
        if _is_abbreviation(full, False):
            if not re.search(r"[a-z]{2}", full):
                continue
        if len(re.sub(r"[^A-Za-z]", "", full)) <= len(re.sub(r"[^A-Za-z]", "", abbr)):
            continue
        return abbr, full
    return None


def _chinese_name(prefix):
    # Stop at sentence boundaries, earlier definitions, citations and line breaks.
    candidate = re.split(r"[，,。；;：:！？!?、()（）\[\]【】\n\r]", prefix)[-1].strip()
    candidate = _CONTEXT.split(candidate)[-1].strip()
    candidate = re.sub(r"^(?:并以|与|或|及|的|一种|一个)+", "", candidate).strip()
    # A conjunction following an English token marks another candidate (IoU或旋转IoU).
    candidate = re.split(r"(?<=[A-Za-z0-9])或", candidate)[-1].strip()
    if not re.search(r"[\u3400-\u9fff]", candidate):
        return ""
    return _clean_space(candidate)


def extract_from_text(text, include_single_letter=False):
    """Extract (abbreviation, English full name, Chinese name) definitions in order.

    Single-letter symbols such as Q/K/V are excluded unless explicitly enabled.
    Chinese names are conservative heuristics and may need manual review.
    """
    results = []
    for match in _PARENTHESES.finditer(text):
        definition = _parse_definition(match.group(1), include_single_letter)
        if definition is None:
            continue
        chinese = _chinese_name(text[:match.start()])
        if chinese:
            results.append((*definition, chinese))
    return results


def _paragraph_texts(doc):
    # XML document order includes nested tables without visiting merged cells twice.
    from docx.oxml.ns import qn

    for paragraph in doc.element.body.iter(qn("w:p")):
        pieces = []
        for node in paragraph.iter():
            if node.tag == qn("w:t"):
                pieces.append(node.text or "")
            elif node.tag in (qn("w:tab"), qn("w:br"), qn("w:cr")):
                pieces.append("\n" if node.tag != qn("w:tab") else "\t")
        yield "".join(pieces)


def extract_abbreviations(docx_path, include_single_letter=False):
    """Return unique (abbreviation, English full name, Chinese name) tuples.

    Full names differing only in case, whitespace or dash style are deduplicated.
    Abbreviation case is significant; different full names remain separate.
    """
    try:
        doc = Document(docx_path)
    except Exception as e:
        print(f"错误：无法打开文档 {docx_path}")
        print(f"详细信息：{e}")
        return []

    abbreviations = []
    seen = set()
    for text in _paragraph_texts(doc):
        for abbr, full, chinese in extract_from_text(text, include_single_letter):
            key = (abbr.translate(_DASHES), full.translate(_DASHES).casefold())
            if key not in seen:
                seen.add(key)
                abbreviations.append((abbr, full, chinese))
    return abbreviations


def save_to_word(abbreviations, template_path, output_path):
    """
    将缩略词保存到Word文档

    Args:
        abbreviations: 缩略词列表
        template_path: 模板文档路径
        output_path: 输出文档路径
    """
    # 加载模板
    doc = Document(template_path)

    # 获取表格（假设模板中只有一个表格）
    table = doc.tables[0]

    # 按英文缩写字母顺序排序
    abbreviations_sorted = sorted(abbreviations, key=lambda x: x[0].upper())

    # 添加数据行
    for abbr, english_full, chinese in abbreviations_sorted:
        row = table.add_row()
        row.cells[0].text = abbr
        row.cells[1].text = english_full
        row.cells[2].text = chinese

    # 保存文档
    doc.save(output_path)


def main():
    if len(sys.argv) < 2:
        print("用法: python extract_abbreviations.py <word文档路径> [模板路径]")
        print("示例: python extract_abbreviations.py 论文.docx")
        print("      python extract_abbreviations.py 论文.docx 缩略词格式.docx")
        sys.exit(1)

    docx_path = sys.argv[1]
    template_path = sys.argv[2] if len(sys.argv) > 2 else "缩略词格式.docx"
    output_path = "缩略词列表.docx"

    print(f"正在提取 {docx_path} 中的缩略词...\n")

    abbreviations = extract_abbreviations(docx_path)

    if not abbreviations:
        print("未找到符合格式的缩略词")
        return

    print(f"共找到 {len(abbreviations)} 个缩略词：\n")
    print("-" * 80)
    print(f"{'序号':<6}{'英文缩写':<15}{'英文全称':<40}{'中文全称':<20}")
    print("-" * 80)

    # 按英文缩写排序显示
    abbreviations_sorted = sorted(abbreviations, key=lambda x: x[0].upper())
    for idx, (abbr, english_full, chinese) in enumerate(abbreviations_sorted, 1):
        print(f"{idx:<6}{abbr:<15}{english_full:<40}{chinese:<20}")

    print("-" * 80)

    # 保存到Word文档
    try:
        save_to_word(abbreviations, template_path, output_path)
        print(f"\n结果已保存到：{output_path}")
    except Exception as e:
        print(f"\n保存Word文档时出错：{e}")
        print("尝试保存为文本文件...")

        # 备用：保存为文本文件
        txt_output = "缩略词列表.txt"
        with open(txt_output, 'w', encoding='utf-8') as f:
            f.write(f"缩略词提取结果（共 {len(abbreviations)} 个）\n")
            f.write("=" * 80 + "\n\n")
            f.write(f"{'英文缩写':<15}{'英文全称':<40}{'中文全称':<20}\n")
            f.write("-" * 80 + "\n")
            for abbr, english_full, chinese in abbreviations_sorted:
                f.write(f"{abbr:<15}{english_full:<40}{chinese:<20}\n")
        print(f"已保存为文本文件：{txt_output}")


if __name__ == "__main__":
    main()
