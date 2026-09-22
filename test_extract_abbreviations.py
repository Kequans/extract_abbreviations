import tempfile
import unittest
from pathlib import Path

from docx import Document

from extract_abbreviations import extract_abbreviations, extract_from_text, save_to_word


class ExtractionTests(unittest.TestCase):
    def test_definition_formats(self):
        cases = [
            ('卷积神经网络（Convolutional Neural Networks，CNN）',
             ('CNN', 'Convolutional Neural Networks', '卷积神经网络')),
            ('本文采用旋转IoU（Rotated IoU, RIoU）', ('RIoU', 'Rotated IoU', '旋转IoU')),
            ('平均精度均值 (mean Average Precision; mAP)',
             ('mAP', 'mean Average Precision', '平均精度均值')),
            ('动态卷积（Convolution-Gram driven Kolmogorov-Arnold Network, C-GKAN）',
             ('C-GKAN', 'Convolution-Gram driven Kolmogorov-Arnold Network', '动态卷积')),
            ('细节增强卷积与组归一化（Detail-Enhanced Convolution with Group Normalization, DEConv-GN）',
             ('DEConv-GN', 'Detail-Enhanced Convolution with Group Normalization', '细节增强卷积与组归一化')),
            ('精确率—召回率（Precision–Recall, PR）', ('PR', 'Precision–Recall', '精确率—召回率')),
            ('目标检测器（DETR，DEtection TRansformer）', ('DETR', 'DEtection TRansformer', '目标检测器')),
            ('三维网络（3D Convolutional Neural Network, 3D-CNN）',
             ('3D-CNN', '3D Convolutional Neural Network', '三维网络')),
            ('干涉 SAR （Interferometric\u00a0SAR, InSAR）', ('InSAR', 'Interferometric SAR', '干涉 SAR')),
            ('并行网络（Parallel Network, PN）', ('PN', 'Parallel Network', '并行网络')),
            ('对比学习（Contrastive Learning, CL）', ('CL', 'Contrastive Learning', '对比学习')),
            ('模型-视图-控制器（Model-View-Controller, MVC）',
             ('MVC', 'Model-View-Controller', '模型-视图-控制器')),
        ]
        for source, expected in cases:
            with self.subTest(source=source):
                self.assertEqual(extract_from_text(source), [expected])

    def test_adjacent_definitions_and_context(self):
        text = ('通过多头注意力（Multi-Head Attention, MHA）与前馈网络（Feed-Forward Network, FFN）。'
                '重叠度为IoU或旋转IoU（Rotated IoU, RIoU）')
        self.assertEqual([r[2] for r in extract_from_text(text)], ['多头注意力', '前馈网络', '旋转IoU'])

    def test_reject_non_definitions(self):
        for text in ['数据（2024, 2025）', '参考文献（Smith, 2024）', '方法（CNN, SAR）',
                     '网络（Convolution Network, cnn）', '术语（中文全称, CNN）',
                     '网络（Convolution Network, CNN, extra）', '网络（Convolution Network）']:
            with self.subTest(text=text):
                self.assertEqual(extract_from_text(text), [])

    def test_single_letter_opt_in(self):
        text = '查询（Query, Q）、键（Key, K）和值（Value, V）'
        self.assertEqual(extract_from_text(text), [])
        self.assertEqual([r[0] for r in extract_from_text(text, True)], ['Q', 'K', 'V'])

    def test_document_order_nested_tables_and_deduplication(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'input.docx'
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            cell = table.cell(0, 0).merge(table.cell(0, 1))
            cell.text = '旋转IoU（Rotated IoU, RIoU）'
            nested = cell.add_table(rows=1, cols=1)
            nested.cell(0, 0).text = '精确率召回率（Precision–Recall, PR）'
            para = doc.add_paragraph('卷积网络（Convolutional ')
            para.add_run('Neural Network, CNN）')
            doc.add_paragraph('重复名称（convolutional  neural\nnetwork, CNN）')
            doc.add_paragraph('重复名称（Precision-Recall, PR）')
            doc.add_paragraph('不同名称（Convolutional New Network, CNN）')
            doc.save(path)
            result = extract_abbreviations(path)
            self.assertEqual([r[0] for r in result], ['RIoU', 'PR', 'CNN', 'CNN'])
            self.assertEqual(result[2][2], '卷积网络')

    def test_word_output_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / 'template.docx'
            output = Path(directory) / 'output.docx'
            doc = Document()
            table = doc.add_table(rows=1, cols=3)
            for cell, label in zip(table.rows[0].cells, ['英文缩写', '英文全称', '中文全称']):
                cell.text = label
            doc.save(template)
            rows = [('SAR', 'Synthetic Aperture Radar', '合成孔径雷达'),
                    ('mAP', 'mean Average Precision', '平均精度均值')]
            save_to_word(rows, template, output)
            actual = Document(output).tables[0]
            self.assertEqual(len(actual.rows), 3)
            self.assertEqual([c.text for c in actual.rows[1].cells], list(rows[1]))


if __name__ == '__main__':
    unittest.main()
